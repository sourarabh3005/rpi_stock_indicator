#!/usr/bin/env python3

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


# ============================================================
# Configuration
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DLOAD_DIR = PROJECT_ROOT / "tmp" / "dload"

SYSTEM_FILE = DLOAD_DIR / "system_info.xlsx"
STOCK_FILE = DLOAD_DIR / "stock_info.xlsx"

MAX_STOCK_ROWS = 100


# ============================================================
# Formatting
# ============================================================

HEADER_FILL = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid"
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF"
)

CENTER = Alignment(
    horizontal="center",
    vertical="center"
)


def format_header(sheet, row=1):
    for cell in sheet[row]:
        if cell.value is not None:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER


def add_type_dropdown(sheet, column, start_row=2, end_row=101):
    """
    Add BUY/SELL dropdown to the Type column.
    """

    validation = DataValidation(
        type="list",
        formula1='"BUY,SELL"',
        allow_blank=True
    )

    validation.error = "Please select BUY or SELL."
    validation.errorTitle = "Invalid Type"

    validation.prompt = "Select BUY or SELL."
    validation.promptTitle = "Stock Type"

    sheet.add_data_validation(validation)

    validation.add(
        f"{column}{start_row}:{column}{end_row}"
    )


# ============================================================
# system_info.xlsx
# ============================================================

def create_system_info():
    wb = Workbook()

    ws = wb.active
    ws.title = "System Info"

    # Required positions used by SystemFields in sheet.py

    ws["A1"] = "Running"

    ws["A2"] = "CPU Temperature"
    ws["A3"] = "Current Date"
    ws["A4"] = "Current Time"

    ws["A6"] = "WiFi SSID"
    ws["A7"] = "IP Address"

    ws["B2"] = 0
    ws["B3"] = ""
    ws["B4"] = ""
    ws["B6"] = ""
    ws["B7"] = ""

    # Stock status
    ws["C1"] = "Buy"
    ws["C2"] = 0

    ws["D1"] = "Sell"
    ws["D2"] = 0

    ws["E1"] = "Critical"
    ws["E2"] = 0

    # Formatting
    for row in ws.iter_rows(
        min_row=1,
        max_row=7,
        min_col=1,
        max_col=5
    ):
        for cell in row:
            cell.alignment = CENTER

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15

    wb.save(SYSTEM_FILE)

    print(f"Created: {SYSTEM_FILE}")


# ============================================================
# stock_info.xlsx
# ============================================================

def create_stock_info():
    wb = Workbook()

    # ========================================================
    # Wishlist
    # ========================================================

    wishlist = wb.active
    wishlist.title = "Wishlist"

    wishlist_headers = [
        "S.No",
        "Ticker",
        "Type",
        "Target",
        "Current"
    ]

    for col, value in enumerate(wishlist_headers, start=1):
        wishlist.cell(
            row=1,
            column=col,
            value=value
        )

    # Add the requested stocks
    wishlist_stocks = [
        "HINDCOPPER",
        "RELIANCE"
    ]

    for index, ticker in enumerate(wishlist_stocks, start=1):
        row = index + 1

        wishlist.cell(row=row, column=1, value=index)
        wishlist.cell(row=row, column=2, value=ticker)

        # Default type
        wishlist.cell(row=row, column=3, value="BUY")

        # Target and Current will be handled by the application
        wishlist.cell(row=row, column=4, value="")
        wishlist.cell(row=row, column=5, value="")

    # Fill serial numbers for remaining rows
    for row in range(
        len(wishlist_stocks) + 2,
        MAX_STOCK_ROWS + 2
    ):
        wishlist.cell(
            row=row,
            column=1,
            value=row - 1
        )

    format_header(wishlist)

    # BUY / SELL dropdown
    add_type_dropdown(
        wishlist,
        "C",
        2,
        MAX_STOCK_ROWS + 1
    )

    wishlist.column_dimensions["A"].width = 10
    wishlist.column_dimensions["B"].width = 18
    wishlist.column_dimensions["C"].width = 12
    wishlist.column_dimensions["D"].width = 15
    wishlist.column_dimensions["E"].width = 15

    # ========================================================
    # Portfolio
    # ========================================================

    portfolio = wb.create_sheet("Portfolio")

    portfolio_headers = [
        "S.No",
        "Ticker",
        "Type",
        "Buy Price",
        "Current",
        "Target",
        "Units",
        "Date",
        "P/L",
        "P/L %"
    ]

    for col, value in enumerate(
        portfolio_headers,
        start=1
    ):
        portfolio.cell(
            row=1,
            column=col,
            value=value
        )

    # Add AMD to portfolio
    portfolio_stocks = [
        "AMD"
    ]

    for index, ticker in enumerate(
        portfolio_stocks,
        start=1
    ):
        row = index + 1

        portfolio.cell(
            row=row,
            column=1,
            value=index
        )

        portfolio.cell(
            row=row,
            column=2,
            value=ticker
        )

        # Default portfolio type
        portfolio.cell(
            row=row,
            column=3,
            value="BUY"
        )

        # Remaining values are intentionally blank.
        #
        # Buy Price
        # Current
        # Target
        # Units
        # Date
        # P/L
        # P/L %

    # Fill serial numbers for remaining rows
    for row in range(
        len(portfolio_stocks) + 2,
        MAX_STOCK_ROWS + 2
    ):
        portfolio.cell(
            row=row,
            column=1,
            value=row - 1
        )

    format_header(portfolio)

    # BUY / SELL dropdown
    add_type_dropdown(
        portfolio,
        "C",
        2,
        MAX_STOCK_ROWS + 1
    )

    portfolio_widths = [
        10, 18, 12, 15, 15,
        15, 12, 15, 15, 15
    ]

    for col, width in enumerate(
        portfolio_widths,
        start=1
    ):
        portfolio.column_dimensions[
            get_column_letter(col)
        ].width = width

    # ========================================================
    # Save
    # ========================================================

    wb.save(STOCK_FILE)

    print(f"Created: {STOCK_FILE}")


# ============================================================
# Main
# ============================================================

def main():

    DLOAD_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Prevent accidental overwrite
    # --------------------------------------------------------

    if SYSTEM_FILE.exists():
        print(
            f"WARNING: {SYSTEM_FILE} already exists"
        )
        print(
            "Delete it manually if you want to recreate it."
        )
    else:
        create_system_info()

    if STOCK_FILE.exists():
        print(
            f"WARNING: {STOCK_FILE} already exists"
        )
        print(
            "Delete it manually if you want to recreate it."
        )
    else:
        create_stock_info()

    print()
    print("Excel initialization completed.")
    print()
    print(f"System file : {SYSTEM_FILE}")
    print(f"Stock file  : {STOCK_FILE}")


if __name__ == "__main__":
    main()